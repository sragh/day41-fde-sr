"""Assemble the evidence pack .docx from the repo's markdown documents."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/artifacts/DecisionStream_Observability_Evidence_Pack.docx"

SECTIONS = [
    ("docs/discovery/01-system-and-business-flow.md", None),
    ("docs/observability/00-baseline-and-gaps.md", None),
    ("docs/observability/02-sli-slo-error-budget.md", None),
    ("docs/observability/03-telemetry-and-correlation-model.md", None),
    ("dashboards/README.md", "dashboards/production_health.png"),
    ("docs/artifacts/incident-01-investigation.md", "docs/artifacts/diagrams/case-1042-waterfall.png"),
    ("runbooks/README.md", None),
    ("docs/artifacts/provenance-and-readiness.md", None),
    ("docs/artifacts/evidence-pack.md", None),
]


def inline(par, text):
    for i, part in enumerate(re.split(r"`", text)):
        run = par.add_run(part)
        if i % 2:
            run.font.name = "Menlo"
            run.font.size = Pt(9)


def add_table(doc, rows):
    rows = [r for r in rows if not re.match(r"^\|[-| ]+\|$", r)]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(cells):
        for j, text in enumerate(row[: len(cells[0])]):
            cell = t.cell(i, j)
            cell.text = ""
            inline(cell.paragraphs[0], text)
            for r in cell.paragraphs[0].runs:
                r.font.size = Pt(9)
                r.bold = r.bold or i == 0
    doc.add_paragraph()


def render(doc, md, image, level_shift):
    lines, i = md.splitlines(), 0
    image_done = False
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            i += 1
            block = []
            while not lines[i].startswith("```"):
                block.append(lines[i]); i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block)); r.font.name = "Menlo"; r.font.size = Pt(9)
        elif ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            add_table(doc, block); continue
        elif ln.startswith("#"):
            n = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln.lstrip("# ").strip(), level=min(n + level_shift, 4))
            if image and n == 2 and not image_done and "Layout" in ln or (image and n == 2 and "Critical path" in ln):
                doc.add_picture(str(ROOT / image), width=Inches(6.3)); image_done = True
        elif re.match(r"^\s*[-*] ", ln):
            inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"^\d+\. ", ln):
            inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", ln))
        elif ln.strip():
            inline(doc.add_paragraph(), ln.strip())
        i += 1
    if image and not image_done:
        doc.add_picture(str(ROOT / image), width=Inches(6.3))


doc = Document()
doc.add_heading("DecisionStream Observability Evidence Pack", 0)
doc.add_paragraph("From black box to explainable production system. Covers tasks 1 to 16 of the AI FDE Observability Challenge.")
for path, image in SECTIONS:
    doc.add_page_break()
    render(doc, (ROOT / path).read_text(), image, 0)
doc.save(OUT)
print(OUT)
